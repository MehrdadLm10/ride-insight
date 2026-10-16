"""Offline parsing, Persian calendar inference and XLSX I/O."""
import hashlib
import re
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from collections import defaultdict

MONTHS = 'فروردین اردیبهشت خرداد تیر مرداد شهریور مهر آبان آذر دی بهمن اسفند'.split()
TRANS = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩', '01234567890123456789')
WEEKDAYS = {'دوشنبه':0, 'سه شنبه':1, 'چهارشنبه':2, 'پنج شنبه':3,
            'جمعه':4, 'شنبه':5, 'یک شنبه':6, 'یکشنبه':6, 'پنجشنبه':3}
TEHRAN = timezone(timedelta(hours=3, minutes=30))

def norm(value):
    return str(value or '').translate(TRANS).replace('\u200c',' ').replace('ي','ی').replace('ك','ک')

def today():
    return datetime.now(TEHRAN).date()

def g2j(g):
    gy, gm, gd = g.year, g.month, g.day
    offsets = [0,31,59,90,120,151,181,212,243,273,304,334]
    gy2 = gy + 1 if gm > 2 else gy
    days = 355666 + 365*gy + (gy2+3)//4 - (gy2+99)//100 + (gy2+399)//400 + gd + offsets[gm-1]
    jy = -1595 + 33*(days//12053); days %= 12053
    jy += 4*(days//1461); days %= 1461
    if days > 365:
        jy += (days-1)//365; days = (days-1)%365
    if days < 186:
        jm, jd = 1+days//31, 1+days%31
    else:
        jm, jd = 7+(days-186)//30, 1+(days-186)%30
    return jy,jm,jd

def j2g(jy,jm,jd):
    jy += 1595
    days = -355668 + 365*jy + (jy//33)*8 + ((jy%33+3)//4) + jd
    days += (jm-1)*31 if jm<7 else (jm-7)*30+186
    gy = 400*(days//146097); days %= 146097
    if days>36524:
        gy += 100*((days-1)//36524); days = (days-1)%36524
        if days>=365: days+=1
    gy += 4*(days//1461); days %= 1461
    if days>365:
        gy += (days-1)//365; days=(days-1)%365
    gd=days+1
    lengths=[31,29 if gy%4==0 and (gy%100!=0 or gy%400==0) else 28,31,30,31,30,31,31,30,31,30,31]
    gm=1
    for n in lengths:
        if gd<=n: break
        gd-=n; gm+=1
    result=date(gy,gm,gd)
    if g2j(result)!=(jy-1595,jm,jd): raise ValueError('تاریخ شمسی نامعتبر')
    return result

def money(text, price_text=''):
    num=r'[0-9]+(?:[,٬][0-9]{3})*'
    for source in (price_text,text):
        source=norm(source)
        m=re.search(rf'(?<![0-9,٬])({num})[ \t\u00a0\u202f]*(تومان|ریال)',source)
        if m: value,unit=m.groups()
        else:
            m=re.search(rf'(تومان|ریال)[ \t\u00a0\u202f]*({num})(?![0-9,٬])',source)
            if not m: continue
            unit,value=m.groups()
        n=int(value.replace(',','').replace('٬',''))
        return n if unit=='تومان' else n/10,unit
    return None,''

def cancelled(text):
    return bool(re.search(r'لغو\s*شده|لغوشده|سفر\s*لغو|کنسل\s*شده|cancelled|canceled',norm(text),re.I))

@dataclass
class Ride:
    key: str
    provider: str
    raw: str
    date_text: str=''
    amount: float|None=None
    status: str=''
    destination: str=''
    origin: str=''
    jalali: str=''
    year: int|None=None
    month: int|None=None
    note: str=''

def date_line(raw, fallback=''):
    # Restrict matching to date-like lines. "مهر" inside names/addresses is not a date.
    lines=[x.strip() for x in norm(raw).splitlines() if x.strip()]
    week=r'یک\s*شنبه|دو\s*شنبه|سه\s*شنبه|چهار\s*شنبه|پنج\s*شنبه|جمعه|شنبه'
    for line in lines:
        if line.startswith('به '): continue
        if re.search(r'\b\d{4}[/\-]\d{1,2}[/\-]\d{1,2}\b',line): return line
        if re.search(r'\d{1,2}\s+(?:'+'|'.join(MONTHS)+r')\b',line) and (line.startswith('سفر') or re.match(week,line)):
            return line
        if line.startswith('سفر') or re.match(week+r'\s*(?:-|شب|ظهر|بعدازظهر|$)',line) or line.startswith(('امروز','دیروز')):
            if re.search(week+r'|امروز|دیروز',line): return line
    # The old date column can contain addresses; inspect only its first segment.
    first=norm(fallback).split('|')[0].strip()
    if re.search(week+r'|امروز|دیروز|\d{4}[/\-]',first): return first
    return ''

def parse_card(card,provider,ordinal=1):
    raw=str(card.get('text') or '')
    amount,_=money(raw,str(card.get('price_text') or ''))
    lines=[x.strip() for x in raw.splitlines() if x.strip()]
    dt=date_line(raw)
    origin='';destination=''
    if provider=='تپسی':
        destination=next((x[3:] for x in lines if x.startswith('به ')),'')
        if 'آخرین سفر شما' in raw:
            for i,x in enumerate(lines):
                if re.search(r'تومان|ریال',x) and i+2<len(lines):
                    if not lines[i+1].startswith('به '): origin=lines[i+1]
                    break
    else:
        normalized=[norm(x) for x in lines]
        idx=next((i for i,x in enumerate(normalized) if x.startswith('سفر')),None)
        if idx is not None:
            addresses=[]
            for x in lines[idx+1:]:
                if x.startswith(('تکرار','برگشت','گزارش')): break
                addresses.append(x)
            if addresses:
                origin=addresses[0]
                if len(addresses)>1: destination=addresses[-1]
    status='لغو شده' if cancelled(raw) else ''
    key=hashlib.sha256((provider+'\0'+raw).encode()).hexdigest()+f':{ordinal}'
    return Ride(key,provider,raw,dt,amount,status,destination,origin)

def infer_dates(rides,anchor):
    """Input newest-to-oldest; never invent an unflagged year or borrow today's year."""
    cursor=anchor
    for ride in rides:
        text=norm(ride.date_text)
        explicit=re.search(r'\b(\d{4})[/\-](\d{1,2})[/\-](\d{1,2})\b',text)
        weekday=None
        for token,val in sorted(WEEKDAYS.items(),key=lambda x:-len(x[0])):
            if token in text: weekday=val;break
        found=None; note=''
        if explicit:
            y,m,d=map(int,explicit.groups())
            try:
                found=j2g(y,m,d) if y<1700 else date(y,m,d)
                if found>cursor: note='تاریخ صریح خارج از ترتیب؛ بررسی شود'
                if weekday is not None and found.weekday()!=weekday: note+='؛ روز هفته ناسازگار'
            except ValueError: note='تاریخ صریح نامعتبر'
        else:
            md=re.search(r'(\d{1,2})\s+('+'|'.join(MONTHS)+r')\b',text)
            if md:
                d=int(md[1]);m=MONTHS.index(md[2])+1
                candidates=[]
                for y in range(g2j(cursor)[0],1300,-1):
                    try: candidate=j2g(y,m,d)
                    except ValueError: continue
                    if candidate<=cursor and (weekday is None or candidate.weekday()==weekday): candidates.append(candidate)
                    if len(candidates)>=2: break
                if candidates:
                    found=candidates[0]
                    note='سال استنباطی از ترتیب و روز هفته؛ فاصله چندساله ممکن است'
                else: note='ماه و روز با ترتیب/روز هفته سازگار نیست'
            elif 'امروز' in text: found=anchor;note='تاریخ نسبی نسبت به زمان استخراج'
            elif 'دیروز' in text: found=anchor-timedelta(days=1);note='تاریخ نسبی نسبت به زمان استخراج'
            elif weekday is not None:
                found=cursor-timedelta(days=(cursor.weekday()-weekday)%7)
                note='تاریخ نسبی استنباطی؛ با زمان استخراج کنترل شود'
            else: note='تاریخ قابل تشخیص نیست'
        if found:
            y,m,d=g2j(found)
            ride.year=y;ride.month=m;ride.jalali=f'{y:04}/{m:02}/{d:02}'
            cursor=min(cursor,found)
        else: ride.year=None;ride.month=None;ride.jalali=''
        ride.note=note
    return rides

def destination_group(text):
    s=norm(text)
    if ('آرژانتین' in s and 'شانزده' in s) or ('ساعی' in s and 'شانزدهم بعد از سام' in s):
        return 'آرژانتین / ساعی، شانزدهم'
    if 'جنت' in s and 'بهمنی' in s: return 'جنت آباد، محدوده بهمنی نژاد'
    if 'تهرانپارس' in s and ('مهدوی' in s or 'مظفری' in s): return 'تهرانپارس، مهدوی زفرقندی / مظفری'
    return re.sub(r'\s+',' ',s).strip(' ،') or 'مقصد نامشخص'

def stats(rides):
    valid=[r for r in rides if not r.status and r.amount is not None]
    months=defaultdict(lambda:[0,0]);years=defaultdict(lambda:[0,0]);dest=defaultdict(lambda:[0,0])
    for r in valid:
        if r.year and r.month:
            months[(r.year,r.month)][0]+=1;months[(r.year,r.month)][1]+=r.amount
            years[r.year][0]+=1;years[r.year][1]+=r.amount
        dest[destination_group(r.destination)][0]+=1;dest[destination_group(r.destination)][1]+=r.amount
    return dict(total=sum(r.amount for r in valid),count=len(valid),cancelled=sum(bool(r.status) for r in rides),
                undated=sum(not r.year for r in valid),missing=sum(r.amount is None and not r.status for r in rides),
                months=dict(months),years=dict(years),dest=dict(dest),valid=valid)

def save_xlsx(path,rides,anchor,reason):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();ws=wb.active;ws.title='سفرها';ws.sheet_view.rightToLeft=True
    ws.append(['ردیف','تاریخ / زمان عین صفحه','مبلغ نمایش داده شده','واحد','وضعیت','لینک کارت','متن کامل سفر','بررسی',
               'سرویس','تاریخ شمسی تحلیلی','سال','ماه','مبدأ','مقصد','شناسه'])
    for i,r in enumerate(rides,1):
        ws.append([i,r.date_text,r.amount,'تومان',r.status,'',r.raw,r.note,r.provider,r.jalali,r.year,r.month,r.origin,r.destination,r.key])
    for row in ws:
        for c in row:
            if isinstance(c.value,str): c.data_type='s'
            c.alignment=Alignment(horizontal='right',vertical='top',wrap_text=True)
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='2563EB')
    for col,width in {'A':8,'B':38,'C':23,'D':10,'E':15,'F':15,'G':70,'H':50,'I':12,'J':20,'K':12,'L':12,'M':50,'N':50,'O':20}.items():ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2,min_col=3,max_col=3):row[0].number_format='#,##0.0'
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    info=wb.create_sheet('راهنما');info.column_dimensions['A'].width=110
    for line in [f'زمان استخراج: {anchor.isoformat()}',f'توقف: {reason}',
                 'سفرهای لغو شده از داشبورد حذف می شوند. وضعیت خالی یعنی بدون علامت لغو؛ تأیید مستقل انجام سفر نیست.',
                 'سال های فاقد سال صریح استنباطی هستند؛ فاصله های چندساله می تواند موجب ابهام شود.',
                 'توقف خودکار در انتهای صفحه کامل بودن تاریخچه را تضمین نمی کند.',
                 'ریال در تحلیل به تومان تبدیل شده است.']:
        info.append([line])
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_name(path.stem+'.tmp.xlsx');wb.save(tmp);tmp.replace(path)

def load_xlsx(path,provider,anchor=None):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=True,read_only=True)
    try:
        if anchor is None and 'راهنما' in wb:
            for row in wb['راهنما'].values:
                if row[0] and 'زمان استخراج:' in str(row[0]):
                    try: anchor=date.fromisoformat(norm(row[0]).split('زمان استخراج:')[1].strip()[:10])
                    except ValueError: pass
        if anchor is None: raise ValueError('زمان استخراج در فایل نیست؛ تاریخ مبنا را در برنامه وارد کنید.')
        sheet=wb['سفرها'] if 'سفرها' in wb else wb.worksheets[0]
        it=iter(sheet.values);headers=[str(v or '').strip() for v in next(it)]
        if 'متن کامل سفر' not in headers: raise ValueError('ستون «متن کامل سفر» در این فایل وجود ندارد.')
        result=[];occurrences=defaultdict(int);saved_dates={}
        for row in it:
            d=dict(zip(headers,row));raw=str(d.get('متن کامل سفر') or '')
            if not raw.strip():continue
            p=str(d.get('سرویس') or provider)
            occurrences[(p,raw)]+=1
            r=parse_card({'text':raw},p,occurrences[(p,raw)])
            if not r.date_text:r.date_text=date_line('',str(d.get('تاریخ / زمان عین صفحه') or ''))
            if cancelled(d.get('وضعیت')):r.status='لغو شده'
            # Prefer reparsing the fare; avoids model digits in old numeric columns.
            if r.amount is None and not r.status and isinstance(d.get('مبلغ نمایش داده شده'),(int,float)):
                r.amount=d['مبلغ نمایش داده شده']/(10 if 'ریال' in str(d.get('واحد')) else 1)
            if d.get('تاریخ شمسی تحلیلی'):
                try:
                    y,m,day=map(int,norm(d['تاریخ شمسی تحلیلی']).split('/'))
                    j2g(y,m,day)
                    saved_dates[r.key]=(y,m,day,str(d.get('بررسی') or 'تاریخ ذخیره شده'))
                except (ValueError,TypeError):pass
            result.append(r)
        for p in set(r.provider for r in result):infer_dates([r for r in result if r.provider==p],anchor)
        for r in result:
            if r.key in saved_dates:
                y,m,d,note=saved_dates[r.key]
                r.year=y;r.month=m;r.jalali=f'{y:04}/{m:02}/{d:02}';r.note=note
        return result,anchor
    finally:wb.close()
